from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

from troca_fontes.analyses import feature_migration_matrices
from troca_fontes.report.workbook import (
    ReferenceReport,
    SheetReport,
    build_workbook,
)


def test_workbook_contains_feature_migration_detail_sheet(tmp_path):
    result = feature_migration_matrices(
        pd.DataFrame({"x": [-1, 0, 1], "x_ste": [0, 1, -1]}),
        features=["x"],
    )["x"]
    reference = ReferenceReport(
        banner="REF ANTIGA",
        match={},
        performance=None,
        drift=None,
        feature_migration={"publico": {"TOTAL": {"x": result}}},
        segments=("TOTAL",),
    )
    report = SheetReport(
        sheet_name="MODEL - SOURCE",
        model_display="MODEL",
        book_display="SOURCE",
        publicos=(),
        referencias=(),
        ref_labels=(),
        volumetrias={"groups": [], "ref_keys": [], "counts": {}, "segments": []},
        notes=(),
        features=("x",),
        references=(reference,),
    )

    output = tmp_path / "report.xlsx"
    build_workbook([report]).save(output)
    loaded = load_workbook(output, data_only=True)

    assert any("MIG VAR" in name for name in loaded.sheetnames)
    detail = loaded[next(name for name in loaded.sheetnames if "MIG VAR" in name)]
    assert "REF ANTIGA" in detail["A1"].value

