"""Workbook writer for summary and migration-matrix outputs.

The original v6 dashboard writer was not present in the recovered source. This
module implements the interface revealed by ``runner.py`` and gives the new
feature-migration analysis an explicit, inspectable Excel representation.
"""

from __future__ import annotations

from dataclasses import dataclass, field
import math
import re
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..analyses import FeatureMigrationMatrix


@dataclass
class ReferenceReport:
    banner: str
    match: Mapping[str, Any]
    performance: Mapping[str, Any] | None
    drift: Mapping[str, Any] | None
    gh_migration: Mapping[str, Any] | None = None
    feature_migration: Mapping[str, Any] | None = None
    segments: Sequence[str] = ()


@dataclass
class SheetReport:
    sheet_name: str
    model_display: str
    book_display: str
    publicos: Sequence[str]
    referencias: Sequence[str]
    ref_labels: Sequence[str]
    volumetrias: Mapping[str, Any]
    notes: Sequence[str]
    features: Sequence[str]
    references: Sequence[ReferenceReport] = field(default_factory=tuple)


_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
_INVALID_TITLE = re.compile(r"[\\/*?:\[\]]")


def _excel(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if hasattr(value, "item"):
        return _excel(value.item())
    return value


def _title(ws: Worksheet, row: int, text: str, *, width: int = 8) -> int:
    ws.cell(row, 1, text)
    ws.cell(row, 1).font = Font(bold=True, color="FFFFFF")
    ws.cell(row, 1).fill = _HEADER_FILL
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    return row + 1


def _table(ws: Worksheet, row: int, headers: Sequence[Any], rows) -> int:
    for col, value in enumerate(headers, 1):
        cell = ws.cell(row, col, _excel(value))
        cell.font = Font(bold=True)
        cell.fill = _SUBHEADER_FILL
    row += 1
    for values in rows:
        for col, value in enumerate(values, 1):
            ws.cell(row, col, _excel(value))
        row += 1
    return row + 1


def _summary_sheet(ws: Worksheet, report: SheetReport) -> None:
    ws.sheet_view.showGridLines = False
    ws["A1"] = f"{report.model_display} — {report.book_display}"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = _HEADER_FILL
    ws.merge_cells("A1:H1")

    row = 3
    row = _title(ws, row, "PÚBLICOS")
    for item in report.publicos:
        ws.cell(row, 1, item)
        row += 1
    row += 1
    row = _title(ws, row, "REFERÊNCIAS")
    for item in report.referencias:
        ws.cell(row, 1, item)
        row += 1

    volumes = report.volumetrias
    row += 1
    row = _title(ws, row, "VOLUMETRIA")
    segments = list(volumes.get("segments", []))
    counts = volumes.get("counts", {})
    groups = volumes.get("groups", [])
    ref_keys = volumes.get("ref_keys", [])
    rows = (
        [label, ref, *[counts.get((key, ref, seg), 0) for seg in segments]]
        for key, label in groups
        for ref in ref_keys
    )
    row = _table(ws, row, ["Público", "Referência", *segments], rows)

    for ref in report.references:
        row = _title(ws, row, ref.banner)
        match_rows = []
        metric_names: list[str] = []
        for population, by_segment in ref.match.items():
            for segment, metrics in by_segment.items():
                if not metric_names:
                    metric_names = list(metrics)
                match_rows.append(
                    [population, segment, *[metrics.get(k) for k in metric_names]]
                )
        row = _table(
            ws,
            row,
            ["Público", "Segmento", *metric_names],
            match_rows,
        )

        if ref.performance:
            metric_names = []
            perf_rows = []
            for segment, metrics in ref.performance.items():
                if not metric_names:
                    metric_names = list(metrics)
                perf_rows.append([segment, *[metrics.get(k) for k in metric_names]])
            row = _table(ws, row, ["Segmento", *metric_names], perf_rows)

        if ref.drift:
            drift_rows = []
            for population, by_segment in ref.drift.items():
                for segment, by_feature in by_segment.items():
                    for feature, metrics in by_feature.items():
                        drift_rows.append(
                            [
                                population,
                                segment,
                                feature,
                                metrics.get("match_6a"),
                                metrics.get("var_pct"),
                                metrics.get("importancia"),
                            ]
                        )
            row = _table(
                ws,
                row,
                [
                    "Público",
                    "Segmento",
                    "Feature",
                    "Match 6 casas",
                    "Variação %",
                    "Importância",
                ],
                drift_rows,
            )

    if report.notes:
        row = _title(ws, row, "OBSERVAÇÕES")
        for note in report.notes:
            ws.cell(row, 1, note)
            row += 1

    ws.freeze_panes = "A2"
    _fit_columns(ws)


def _matrix_table(ws: Worksheet, row: int, result: FeatureMigrationMatrix) -> int:
    matrix = result.matrix
    ws.cell(row, 1, "Limites antigos (D00..D10)")
    ws.cell(row, 2, "; ".join(f"{x:.8g}" for x in result.old_edges))
    row += 1
    ws.cell(row, 1, "Limites novos (D00..D10)")
    ws.cell(row, 2, "; ".join(f"{x:.8g}" for x in result.new_edges))
    row += 1
    return _table(
        ws,
        row,
        [matrix.index.name, *matrix.columns.tolist()],
        ([idx, *matrix.loc[idx].tolist()] for idx in matrix.index),
    )


def _feature_migration_sheet(ws: Worksheet, report: SheetReport) -> None:
    row = 1
    for ref in report.references:
        if not ref.feature_migration:
            continue
        for population, by_segment in ref.feature_migration.items():
            for segment, by_feature in by_segment.items():
                for feature, result in by_feature.items():
                    row = _title(
                        ws,
                        row,
                        f"{ref.banner} | {population} | {segment} | {feature}",
                        width=13,
                    )
                    row = _matrix_table(ws, row, result)
    ws.freeze_panes = "A2"
    _fit_columns(ws, max_width=45)


def _gh_migration_sheet(ws: Worksheet, report: SheetReport) -> None:
    row = 1
    for ref in report.references:
        if not ref.gh_migration:
            continue
        for population, by_segment in ref.gh_migration.items():
            for segment, matrix in by_segment.items():
                row = _title(
                    ws,
                    row,
                    f"{ref.banner} | {population} | {segment}",
                    width=max(8, len(matrix.columns) + 1),
                )
                row = _table(
                    ws,
                    row,
                    [matrix.index.name or "GH antiga", *matrix.columns.tolist()],
                    ([idx, *matrix.loc[idx].tolist()] for idx in matrix.index),
                )
    ws.freeze_panes = "A2"
    _fit_columns(ws)


def _fit_columns(ws: Worksheet, *, max_width: int = 60) -> None:
    for column_cells in ws.columns:
        width = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=8,
        )
        column_letter = get_column_letter(column_cells[0].column)
        ws.column_dimensions[column_letter].width = min(
            max(width + 2, 10), max_width
        )
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def _unique_title(base: str, used: set[str]) -> str:
    clean = _INVALID_TITLE.sub("-", base).strip() or "Relatório"
    clean = clean[:31]
    candidate = clean
    counter = 2
    while candidate.casefold() in used:
        suffix = f" ({counter})"
        candidate = f"{clean[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(candidate.casefold())
    return candidate


def build_workbook(reports: Sequence[SheetReport]) -> Workbook:
    """Build one summary plus GH/feature-migration detail tabs per report."""

    workbook = Workbook()
    workbook.remove(workbook.active)
    used: set[str] = set()
    for report in reports:
        summary = workbook.create_sheet(_unique_title(report.sheet_name, used))
        _summary_sheet(summary, report)
        if any(ref.gh_migration for ref in report.references):
            gh = workbook.create_sheet(
                _unique_title(f"{report.sheet_name} - MIG GH", used)
            )
            _gh_migration_sheet(gh, report)
        if any(ref.feature_migration for ref in report.references):
            feature = workbook.create_sheet(
                _unique_title(f"{report.sheet_name} - MIG VAR", used)
            )
            _feature_migration_sheet(feature, report)
    if not reports:
        workbook.create_sheet("Relatório")
    return workbook
